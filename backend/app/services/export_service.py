import io
import logging
from datetime import datetime
from xml.sax.saxutils import escape

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from docx import Document
from docx.shared import Inches, Pt, RGBColor

logger = logging.getLogger(__name__)


def _score_text(report_data: dict) -> str:
    score = report_data.get("health_score")
    if score is None:
        return "Not available (no results could be read from this document)"
    return f"{score}/100"


def _caveats(report_data: dict) -> list:
    """Lines that must travel with the export so a reader sees the same limits
    the app shows on screen."""
    lines = []
    if report_data.get("health_score_basis"):
        lines.append(f"Score basis: {report_data['health_score_basis']}")
    if report_data.get("analysis_mode") == "rule_based":
        lines.append(
            "Produced by the rule-based reader: every value, unit and reference range below "
            "was read from the uploaded document. No AI model was used and nothing was inferred."
        )
    if report_data.get("llm_notice"):
        lines.append(report_data["llm_notice"])
    lines.extend(report_data.get("warnings") or [])
    unparsed = report_data.get("unparsed_lines") or []
    if unparsed:
        lines.append(
            f"{len(unparsed)} line(s) of the document could not be read as results and are not "
            "shown below. Check them against the original report."
        )
    return lines


def generate_report_pdf(report_data: dict) -> bytes:
    """Generate a clean, printable PDF of the analysis report."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40
    )
    
    styles = getSampleStyleSheet()
    
    # Custom Styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=22,
        textColor=colors.HexColor('#0F766E'),  # Deep Teal
        spaceAfter=15
    )
    section_title_style = ParagraphStyle(
        'SectionTitle',
        parent=styles['Heading2'],
        fontSize=14,
        textColor=colors.HexColor('#0F766E'),  # Teal
        spaceBefore=12,
        spaceAfter=8
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#1F2937'),
        spaceAfter=6
    )
    caveat_style = ParagraphStyle(
        'DocCaveat',
        parent=styles['Normal'],
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#B45309'),  # Amber
        spaceAfter=4
    )
    disclaimer_style = ParagraphStyle(
        'DocDisclaimer',
        parent=styles['Normal'],
        fontSize=8,
        leading=11,
        textColor=colors.HexColor('#DC2626'),  # Red warning
        spaceBefore=15
    )
    
    story = []
    
    # Title
    story.append(Paragraph(f"MediSense AI - {escape(str(report_data.get('report_type', 'Medical Analysis')))}", title_style))
    story.append(Paragraph(f"<b>Patient Name:</b> {escape(str(report_data.get('patient_name', 'Patient')))}", body_style))
    story.append(Paragraph(f"<b>Report Date:</b> {escape(str(report_data.get('created_at', datetime.now().strftime('%Y-%m-%d %H:%M'))))}", body_style))
    story.append(Paragraph(f"<b>Overall Health Score:</b> {escape(_score_text(report_data))}", body_style))
    story.append(Spacer(1, 10))
    
    # Summary
    story.append(Paragraph("Overall Summary", section_title_style))
    story.append(Paragraph(escape(str(report_data.get("summary", ""))), body_style))
    story.append(Spacer(1, 10))

    caveats = _caveats(report_data)
    if caveats:
        story.append(Paragraph("How to read this document", section_title_style))
        for line in caveats:
            story.append(Paragraph(f"• {escape(str(line))}", caveat_style))
        story.append(Spacer(1, 10))
    
    # Biomarkers Table
    story.append(Paragraph("Results Read From This Report", section_title_style))
    
    table_data = [["Test", "Value", "Reference Range", "Status"]]
    for bm in report_data.get("biomarkers", []):
        table_data.append([
            escape(str(bm.get("name", ""))),
            escape(f"{bm.get('value', '')} {bm.get('unit', '')}".strip()),
            escape(str(bm.get("reference_range", "") or "Not printed")),
            escape(str(bm.get("status", "")))
        ])
    if len(table_data) == 1:
        table_data.append(["No results could be read", "", "", ""])
        
    t = Table(table_data, colWidths=[180, 100, 140, 90])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0F766E')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0,0), (-1,0), 6),
        ('BACKGROUND', (0,1), (-1,-1), colors.HexColor('#F3F4F6')),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E5E7EB')),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F9FAFB')]),
        ('FONTSIZE', (0,0), (-1,-1), 9),
    ]))
    story.append(t)
    story.append(Spacer(1, 15))

    concerns = report_data.get("health_concerns") or {}
    risks = concerns.get("potential_risks") or []
    if concerns.get("summary") or risks:
        story.append(Paragraph("Observed Findings", section_title_style))
        if concerns.get("summary"):
            story.append(Paragraph(escape(str(concerns["summary"])), body_style))
        for risk in risks:
            story.append(Paragraph(f"• {escape(str(risk))}", body_style))
        story.append(Spacer(1, 10))
    
    # Recommendations
    story.append(Paragraph("Educational Guidance", section_title_style))
    rec = report_data.get("recommendations", {})
    
    story.append(Paragraph("<b>Dietary Changes:</b>", body_style))
    for d in rec.get("diet", []):
        story.append(Paragraph(f"• {escape(str(d))}", body_style))
        
    story.append(Paragraph("<b>Lifestyle Changes:</b>", body_style))
    for l in rec.get("lifestyle", []):
        story.append(Paragraph(f"• {escape(str(l))}", body_style))
        
    story.append(Paragraph("<b>Recommended Specialists:</b>", body_style))
    story.append(Paragraph(escape(", ".join(str(s) for s in rec.get("specialists", []))), body_style))
    story.append(Paragraph(
        "This document does not prescribe medicines or change doses. Only a clinician who knows the "
        "patient's history can decide treatment.",
        body_style
    ))
    
    # Disclaimer
    story.append(Spacer(1, 15))
    story.append(Paragraph(
        "<b>MEDICAL DISCLAIMER:</b> MediSense AI is an educational analysis service and is NOT a substitute "
        "for professional medical advice, diagnosis, or treatment. Always consult with a qualified physician or "
        "healthcare specialist regarding any medical concerns.", 
        disclaimer_style
    ))
    
    doc.build(story)
    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes

def generate_report_docx(report_data: dict) -> bytes:
    """Generate a clean, formatted Word (.docx) document of the AI analysis."""
    doc = Document()
    
    # Document Style Settings
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(10.5)
    
    # Title
    title = doc.add_paragraph()
    title_run = title.add_run(f"MediSense AI Analysis - {report_data.get('report_type', 'Medical Report')}")
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(30, 58, 138)  # Deep Navy Blue
    
    # Patient Info
    info_p = doc.add_paragraph()
    info_p.add_run("Patient Name: ").bold = True
    info_p.add_run(f"{report_data.get('patient_name', 'Patient')}\n")
    info_p.add_run("Report Date: ").bold = True
    info_p.add_run(f"{report_data.get('created_at', datetime.now().strftime('%Y-%m-%d %H:%M'))}\n")
    info_p.add_run("General Health Score: ").bold = True
    info_p.add_run(_score_text(report_data))
    
    doc.add_paragraph("").paragraph_format.space_after = Pt(6)
    
    # Summary
    h_sum = doc.add_paragraph()
    h_sum_run = h_sum.add_run("Overall Summary")
    h_sum_run.font.size = Pt(14)
    h_sum_run.font.bold = True
    h_sum_run.font.color.rgb = RGBColor(15, 118, 110)  # Teal
    doc.add_paragraph(report_data.get("summary", ""))

    caveats = _caveats(report_data)
    if caveats:
        h_cav = doc.add_paragraph()
        h_cav_run = h_cav.add_run("How to read this document")
        h_cav_run.font.size = Pt(14)
        h_cav_run.font.bold = True
        h_cav_run.font.color.rgb = RGBColor(180, 83, 9)  # Amber
        for line in caveats:
            bullet = doc.add_paragraph(str(line), style='List Bullet')
            for run in bullet.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(180, 83, 9)
    
    # Biomarkers Table
    h_table = doc.add_paragraph()
    h_table_run = h_table.add_run("Results Read From This Report")
    h_table_run.font.size = Pt(14)
    h_table_run.font.bold = True
    h_table_run.font.color.rgb = RGBColor(15, 118, 110)
    
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Test'
    hdr_cells[1].text = 'Value'
    hdr_cells[2].text = 'Reference Range'
    hdr_cells[3].text = 'Status'
    
    biomarkers = report_data.get("biomarkers", [])
    for bm in biomarkers:
        row_cells = table.add_row().cells
        row_cells[0].text = str(bm.get("name", ""))
        row_cells[1].text = f"{bm.get('value', '')} {bm.get('unit', '')}".strip()
        row_cells[2].text = str(bm.get("reference_range", "") or "Not printed")
        row_cells[3].text = str(bm.get("status", ""))
    if not biomarkers:
        table.add_row().cells[0].text = "No results could be read from this document"

    doc.add_paragraph("").paragraph_format.space_before = Pt(12)

    concerns = report_data.get("health_concerns") or {}
    risks = concerns.get("potential_risks") or []
    if concerns.get("summary") or risks:
        h_con = doc.add_paragraph()
        h_con_run = h_con.add_run("Observed Findings")
        h_con_run.font.size = Pt(14)
        h_con_run.font.bold = True
        h_con_run.font.color.rgb = RGBColor(15, 118, 110)
        if concerns.get("summary"):
            doc.add_paragraph(str(concerns["summary"]))
        for risk in risks:
            doc.add_paragraph(str(risk), style='List Bullet')
    
    # Recommendations
    h_rec = doc.add_paragraph()
    h_rec_run = h_rec.add_run("Educational Guidance")
    h_rec_run.font.size = Pt(14)
    h_rec_run.font.bold = True
    h_rec_run.font.color.rgb = RGBColor(15, 118, 110)
    
    rec = report_data.get("recommendations", {})
    
    doc.add_paragraph("Dietary Suggestions:").bold = True
    for d in rec.get("diet", []):
        doc.add_paragraph(str(d), style='List Bullet')
        
    doc.add_paragraph("Lifestyle Modifications:").bold = True
    for l in rec.get("lifestyle", []):
        doc.add_paragraph(str(l), style='List Bullet')
        
    doc.add_paragraph("Specialists to Consult:").bold = True
    doc.add_paragraph(", ".join(str(s) for s in rec.get("specialists", [])))
    doc.add_paragraph(
        "This document does not prescribe medicines or change doses. Only a clinician who knows the "
        "patient's history can decide treatment."
    )
    
    # Disclaimer
    doc.add_paragraph("").paragraph_format.space_before = Pt(12)
    disc_p = doc.add_paragraph()
    disc_run = disc_p.add_run(
        "MEDICAL DISCLAIMER: MediSense AI is an educational analysis service and is NOT a substitute "
        "for professional medical advice, diagnosis, or treatment. Always consult with a qualified physician or "
        "healthcare specialist regarding any medical concerns."
    )
    disc_run.font.size = Pt(8.5)
    disc_run.italic = True
    disc_run.font.color.rgb = RGBColor(220, 38, 38)
    
    buffer = io.BytesIO()
    doc.save(buffer)
    docx_bytes = buffer.getvalue()
    buffer.close()
    return docx_bytes
