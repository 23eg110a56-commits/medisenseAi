"""Tests that exported documents carry the same limits the UI shows on screen.

An exported PDF leaves the application, so it cannot rely on a banner that is no
longer visible: every caveat has to be printed into the file itself.
"""

import io

import pymupdf
from docx import Document

from app.services.export_service import generate_report_docx, generate_report_pdf
from app.services.report_parser import analyze_text

LIPID_TEXT = """LIPID PROFILE
Total Cholesterol      248.0    mg/dL    < 200.0     High
HDL Cholesterol        34.5     mg/dL    > 40.0      Low
LDL Cholesterol        92.0     mg/dL    < 100.0     Normal
Triglycerides          118.0    mg/dL    < 150.0     Normal
"""


def _analysis(**overrides):
    data = analyze_text(LIPID_TEXT)
    data.update({
        "patient_name": "Jane Doe",
        "created_at": "2026-10-09 10:30",
        "report_type": "Lipid Profile",
    })
    data.update(overrides)
    return data


def _normalise(text: str) -> str:
    """Collapse line wraps so multi-word phrases can be asserted on."""
    return " ".join(text.split())


def _pdf_text(pdf_bytes: bytes) -> str:
    with pymupdf.open(stream=pdf_bytes, filetype="pdf") as doc:
        return _normalise("\n".join(page.get_text() for page in doc))


def _docx_text(docx_bytes: bytes) -> str:
    doc = Document(io.BytesIO(docx_bytes))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return _normalise("\n".join(parts))


def _both_formats(report: dict) -> dict:
    return {
        "pdf": _pdf_text(generate_report_pdf(report)),
        "docx": _docx_text(generate_report_docx(report)),
    }


def test_exports_are_non_empty_binary_documents():
    report = _analysis()
    pdf = generate_report_pdf(report)
    docx = generate_report_docx(report)
    assert pdf.startswith(b"%PDF")
    assert len(pdf) > 1000
    assert docx.startswith(b"PK")
    assert len(docx) > 1000


def test_printed_values_units_and_ranges_survive_the_export():
    for text in _both_formats(_analysis()).values():
        for expected in ("248.0", "mg/dL", "< 200.0", "34.5", "92.0", "Total Cholesterol"):
            assert expected in text


def test_no_literal_none_is_printed_as_a_score():
    """Regression: a null health score used to render as "None/100"."""
    report = _analysis(health_score=None, biomarkers=[])
    for label, text in _both_formats(report).items():
        assert "None/100" not in text, label
        assert "None" not in text, label
        assert "Not available" in text, label


def test_disclaimer_appears_in_both_formats():
    for label, text in _both_formats(_analysis()).items():
        assert "NOT a substitute" in text, label
        assert "qualified" in text.lower(), label


def test_rule_based_provenance_is_disclosed():
    for label, text in _both_formats(_analysis(analysis_mode="rule_based")).items():
        assert "rule-based" in text.lower(), label
        assert "No AI model was used" in text, label


def test_llm_notice_is_carried_into_the_export():
    notice = "The AI model could not be reached, so this analysis was produced by the built-in parser."
    for label, text in _both_formats(_analysis(llm_notice=notice)).items():
        assert notice in text, label


def test_urgent_warnings_are_printed_not_dropped():
    report = analyze_text("Hemoglobin    4.1    g/dL    13.5 - 17.5    Low\n")
    report.update({"patient_name": "Jane Doe", "created_at": "2026-10-09 10:30"})
    assert report["warnings"]
    for label, text in _both_formats(report).items():
        assert "Seek prompt evaluation" in text, label


def test_observed_findings_section_is_present_in_both_formats():
    for label, text in _both_formats(_analysis()).items():
        assert "Observed" in text, label
        assert "not a diagnosis" in text, label


def test_unparsed_line_count_is_disclosed():
    report = _analysis(unparsed_lines=["Erythrocyte Sed Rate 22 mm first hour Westergren"])
    for label, text in _both_formats(report).items():
        assert "could not be read as results" in text, label


def test_empty_analysis_says_so_instead_of_showing_a_blank_table():
    report = _analysis(biomarkers=[], health_score=None)
    for label, text in _both_formats(report).items():
        assert "No results could be read" in text, label


def test_no_prescription_language_in_the_export():
    for label, text in _both_formats(_analysis()).items():
        assert "does not prescribe" in text.lower(), label


def test_special_characters_do_not_corrupt_the_documents():
    report = _analysis(patient_name="O'Brien & < Sons > \"Ltd\"")
    report["biomarkers"] = [
        {
            "name": "Total Cholesterol <fasting> & LDL",
            "value": "248.0",
            "unit": "mg/dL",
            "reference_range": "< 200.0",
            "status": "High",
        }
    ]
    pdf = generate_report_pdf(report)
    docx = generate_report_docx(report)
    assert pdf.startswith(b"%PDF")
    assert docx.startswith(b"PK")
    assert "O'Brien" in _pdf_text(pdf)
    assert "O'Brien" in _docx_text(docx)


def test_export_of_a_minimal_report_dict_does_not_crash():
    """A partially populated record must still produce a readable document."""
    assert generate_report_pdf({}).startswith(b"%PDF")
    assert generate_report_docx({}).startswith(b"PK")
